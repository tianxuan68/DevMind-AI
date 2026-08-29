from __future__ import annotations

import json

import fitz
import pytest
from docx import Document as DocxDocument

from internal_kb_qa.document_loaders import load
from internal_kb_qa.document_loaders.confluence_loader import ConfluenceLoader
from internal_kb_qa.document_loaders.markdown_loader import MarkdownLoader
from internal_kb_qa.document_loaders.openapi_loader import OpenAPILoader
from internal_kb_qa.document_loaders.pdf_loader import PDFLoader
from internal_kb_qa.document_loaders.word_loader import WordLoader
from internal_kb_qa.text_splitters.recursive_splitter import RecursiveSplitter


def test_markdown_loader_preserves_code_table_and_metadata(tmp_path):
    path = tmp_path / "deployment_runbook.md"
    path.write_text(
        "# Deployment\n\n| key | value |\n| --- | --- |\n| port | 8003 |\n\n```bash\nuv run python app.py\n```\n",
        encoding="utf-8",
    )

    documents = MarkdownLoader().load(str(path))

    assert len(documents) == 1
    assert "uv run python app.py" in documents[0].page_content
    assert "| port | 8003 |" in documents[0].page_content
    assert documents[0].metadata == {
        "source": path.name,
        "page": 1,
        "doc_type": "runbook",
    }


def test_pdf_loader_returns_one_document_per_non_empty_page(tmp_path):
    path = tmp_path / "incident_report.pdf"
    pdf = fitz.open()
    first_page = pdf.new_page()
    first_page.insert_text((72, 72), "Incident report page one")
    second_page = pdf.new_page()
    second_page.insert_text((72, 72), "Incident report page two")
    pdf.save(path)
    pdf.close()

    documents = PDFLoader().load(str(path))

    assert [document.metadata["page"] for document in documents] == [1, 2]
    assert all(document.metadata["doc_type"] == "report" for document in documents)
    assert "page one" in documents[0].page_content


def test_word_loader_extracts_paragraphs_and_tables(tmp_path):
    path = tmp_path / "faq.docx"
    word = DocxDocument()
    word.add_heading("Dependency FAQ", level=1)
    word.add_paragraph("Install the dependencies with uv sync.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Command"
    table.cell(0, 1).text = "Purpose"
    table.cell(1, 0).text = "uv sync"
    table.cell(1, 1).text = "Install environment"
    word.save(path)

    documents = WordLoader().load(str(path))

    assert len(documents) == 1
    assert "Install the dependencies with uv sync." in documents[0].page_content
    assert "| uv sync | Install environment |" in documents[0].page_content
    assert documents[0].metadata["doc_type"] == "faq"


def test_confluence_loader_removes_navigation_and_keeps_content(tmp_path):
    path = tmp_path / "wiki.html"
    path.write_text(
        """
        <html><body>
          <nav>Navigation noise</nav>
          <div id="main-content">
            <h1>Service guide</h1>
            <p>Use the health endpoint to check service status.</p>
            <pre>curl http://localhost:8003/health</pre>
            <table><tr><th>Code</th><th>Meaning</th></tr><tr><td>200</td><td>Healthy</td></tr></table>
          </div>
          <script>should not appear</script>
        </body></html>
        """,
        encoding="utf-8",
    )

    documents = ConfluenceLoader().load(str(path))

    assert len(documents) == 1
    content = documents[0].page_content
    assert "Navigation noise" not in content
    assert "should not appear" not in content
    assert "curl http://localhost:8003/health" in content
    assert "Healthy" in content
    assert documents[0].metadata["doc_type"] == "wiki"


def test_openapi_loader_formats_paths_parameters_and_responses(tmp_path):
    path = tmp_path / "openapi.json"
    path.write_text(
        json.dumps(
            {
                "openapi": "3.0.3",
                "info": {"title": "Internal API", "version": "1.0.0"},
                "paths": {
                    "/health": {
                        "get": {
                            "operationId": "health",
                            "summary": "Health check",
                            "responses": {"200": {"description": "Healthy"}},
                        }
                    },
                    "/users/{user_id}": {
                        "parameters": [
                            {
                                "name": "user_id",
                                "in": "path",
                                "required": True,
                                "schema": {"type": "string"},
                            }
                        ],
                        "get": {
                            "description": "Get one user",
                            "responses": {"404": {"description": "Not found"}},
                        },
                    },
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    documents = OpenAPILoader().load(str(path))

    assert len(documents) == 1
    content = documents[0].page_content
    assert "# Internal API" in content
    assert "### GET /health" in content
    assert "### GET /users/{user_id}" in content
    assert "user_id (path, required)" in content
    assert "404: Not found" in content
    assert documents[0].metadata["doc_type"] == "api"


def test_load_dispatches_by_extension(tmp_path):
    path = tmp_path / "notes.MARKDOWN"
    path.write_text("# Notes", encoding="utf-8")

    documents = load(str(path))

    assert documents[0].page_content == "# Notes"


def test_load_rejects_missing_and_unsupported_files(tmp_path):
    with pytest.raises(FileNotFoundError, match="does not exist"):
        load(str(tmp_path / "missing.md"))

    path = tmp_path / "notes.txt"
    path.write_text("plain text", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported document extension"):
        load(str(path))


def test_openapi_loader_reports_invalid_json(tmp_path):
    path = tmp_path / "broken.json"
    path.write_text("{broken", encoding="utf-8")

    with pytest.raises(json.JSONDecodeError):
        OpenAPILoader().load(str(path))


def test_recursive_splitter_preserves_metadata():
    from langchain_core.documents import Document

    chunks = RecursiveSplitter(chunk_size=24, chunk_overlap=4).split(
        [
            Document(
                page_content="第一段内容很长，需要切分。\n\n第二段也需要保留来源信息。",
                metadata={"page": 3},
            )
        ]
    )

    assert len(chunks) >= 2
    assert all(chunk.metadata["page"] == 3 for chunk in chunks)
